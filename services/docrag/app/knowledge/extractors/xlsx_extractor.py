"""xlsx extraction: each sheet becomes one PageResult-shaped table block,
reusing the same downstream reconstruction path as PDFs so xlsx and PDF
inputs converge on the same KnowledgeDocument schema."""

from __future__ import annotations

import openpyxl

from app.knowledge.extractors.pdf_extractor import ExtractedTable, PageResult


def extract_xlsx(path: str) -> list[PageResult]:
    wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
    results: list[PageResult] = []

    for i, sheet_name in enumerate(wb.sheetnames):
        ws = wb[sheet_name]
        rows = []
        for row in ws.iter_rows(values_only=True):
            if all(c is None for c in row):
                continue
            rows.append(["" if c is None else str(c) for c in row])

        if not rows:
            results.append(
                PageResult(page_number=i + 1, markdown=f"# {sheet_name}\n\n[empty sheet]", tables=[],
                           engine="local_text_layer")
            )
            continue

        markdown = f"# {sheet_name}\n"
        # Spreadsheet cells have no PDF-point geometry, so no bbox.
        results.append(
            PageResult(
                page_number=i + 1,
                markdown=markdown,
                tables=[ExtractedTable(rows=rows)],
                engine="local_text_layer",
            )
        )

    wb.close()
    return results


def extract_xls(path: str) -> list[PageResult]:
    """Legacy Excel 97–2003 workbooks (CMA data still arrives as .xls): the same
    one-table-per-sheet shape as extract_xlsx, so both converge downstream."""
    import xlrd

    book = xlrd.open_workbook(path, on_demand=True)
    results: list[PageResult] = []
    try:
        for i in range(book.nsheets):
            sheet = book.sheet_by_index(i)
            rows = []
            for r in range(sheet.nrows):
                values = sheet.row_values(r)
                if all(v in ("", None) for v in values):
                    continue
                rows.append([_cell_text(v) for v in values])
            if not rows:
                results.append(PageResult(page_number=i + 1,
                                          markdown=f"# {sheet.name}\n\n[empty sheet]",
                                          tables=[], engine="local_text_layer"))
                continue
            results.append(PageResult(page_number=i + 1, markdown=f"# {sheet.name}\n",
                                      tables=[ExtractedTable(rows=rows)],
                                      engine="local_text_layer"))
            book.unload_sheet(i)
    finally:
        book.release_resources()
    return results


def _cell_text(value: object) -> str:
    # xlrd returns every number as float: 1200.0 reads better as 1200.
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return "" if value is None else str(value)
